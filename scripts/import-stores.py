#!/usr/bin/env python3
"""Convert the provided Nintoki store workbook into the static site's source data."""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/home/ubuntu/upload/YAKINIKUEN_忍鬨_全店舗一覧_2026-10-23.xlsx")
DESTINATION = Path(__file__).resolve().parents[1] / "src" / "data" / "stores.json"

SLUGS = [
    "higashishinsaibashi",
    "higashishinsaibashi-apa",
    "nihonbashi",
    "sakaisujihommachi",
    "noda",
    "tondabayashi",
    "osaka-umeda",
    "umeda-higashidori",
]

HOURS = {
    "higashishinsaibashi": [{"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], "opens": "16:00", "closes": "00:00"}],
    "higashishinsaibashi-apa": [
        {"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"], "opens": "07:00", "closes": "10:00"},
        {"dayOfWeek": ["Saturday", "Sunday"], "opens": "07:30", "closes": "10:00"},
        {"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], "opens": "15:00", "closes": "00:00"},
    ],
    "nihonbashi": [
        {"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday"], "opens": "17:00", "closes": "00:00"},
        {"dayOfWeek": ["Friday", "Saturday", "Sunday"], "opens": "16:00", "closes": "00:00"},
    ],
    "sakaisujihommachi": [{"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], "opens": "17:00", "closes": "00:00"}],
    "noda": [{"dayOfWeek": ["Monday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], "opens": "17:00", "closes": "00:00"}],
    "tondabayashi": [{"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], "opens": "17:00", "closes": "23:00"}],
    "osaka-umeda": [
        {"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday"], "opens": "17:00", "closes": "00:00"},
        {"dayOfWeek": ["Friday", "Saturday"], "opens": "16:00", "closes": "00:00"},
        {"dayOfWeek": ["Sunday"], "opens": "16:00", "closes": "23:00"},
    ],
    "umeda-higashidori": [{"dayOfWeek": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], "opens": "18:00", "closes": "06:00"}],
}


def intl_phone(phone: str) -> str:
    digits = re.sub(r"\D", "", phone)
    return f"+81{digits[1:]}" if digits.startswith("0") else phone


def locality(address: str) -> str:
    match = re.search(r"大阪府([^0-9]+?(?:市|区))", address)
    return match.group(1) if match else "大阪府"


def main() -> None:
    wb = load_workbook(SOURCE, data_only=True, read_only=True)
    ws = wb["店舗一覧"]
    stores = []
    menu_by_name: dict[str, list[dict]] = defaultdict(list)

    for row in ws.iter_rows(min_row=28, values_only=True):
        name, status, category, item, price, source, page, note = row[:8]
        if not name or name.startswith("メニュー記載基準"):
            continue
        if status != "メニュー掲載" or not item or item == "メニュー資料未提供":
            continue
        menu_by_name[str(name)].append({
            "category": str(category),
            "name": str(item),
            "price": int(price) if isinstance(price, (int, float)) else None,
            "note": str(note) if note else "",
            "source": str(source) if source else "",
            "page": int(page) if isinstance(page, (int, float)) else None,
        })

    for index, row in enumerate(ws.iter_rows(min_row=4, max_row=11, values_only=True)):
        _, name, postal, address, phone, opening, closed, access, _, _, menu_source, item_count, verified = row[:13]
        slug = SLUGS[index]
        stores.append({
            "slug": slug,
            "name": str(name),
            "postalCode": str(postal),
            "address": str(address),
            "addressLocality": locality(str(address)),
            "addressRegion": "大阪府",
            "telephone": str(phone),
            "telephoneInternational": intl_phone(str(phone)),
            "openingHoursText": str(opening).replace("\n", "<br>"),
            "openingHoursSchema": HOURS[slug],
            "closedText": str(closed),
            "access": str(access),
            "menuSource": str(menu_source),
            "menuStatus": "available" if int(item_count or 0) else "pending",
            "menuItemCount": int(item_count or 0),
            "verifiedDate": str(verified),
            "menu": menu_by_name[str(name)],
        })

    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    DESTINATION.write_text(json.dumps({"updated": "2026-10-23", "stores": stores}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {DESTINATION} with {len(stores)} stores and {sum(len(s['menu']) for s in stores)} menu items.")


if __name__ == "__main__":
    main()
